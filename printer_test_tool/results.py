"""시험 결과 저장 (프린터별) / CSV 내보내기 / 계획표 엑셀에 결과 반영."""
import csv
import datetime
import json
import os

from engine import BASE_DIR
from testcases import TC

RESULT_DIR = os.path.join(BASE_DIR, "results")
RESULT_FILE = os.path.join(RESULT_DIR, "results.json")

# 계획표 '체크리스트' 시트 열 위치 (A=1)
COL_ID, COL_RESULT, COL_TESTER, COL_DATE, COL_FW, COL_DEFECT, COL_NOTE = 2, 7, 8, 9, 10, 11, 12
EMPTY = {"result": "미실시", "tester": "", "date": "", "fw": "", "defect": "", "note": ""}


def _safe(name):
    return "".join("_" if c in '\\/:*?"<>|' else c for c in name)


class Results:
    """data = {프린터이름: {TC ID: {result, tester, date, fw, defect, note}}}"""

    def __init__(self, path=RESULT_FILE):
        self.path = os.path.abspath(path)
        self.data = {}
        if os.path.exists(self.path):
            try:
                with open(self.path, encoding="utf-8") as f:
                    raw = json.load(f)
                # 예전(프린터 구분 없는) 형식이면 '기본' 프린터로 옮긴다
                if raw and all(isinstance(v, dict) and "result" in v for v in raw.values()):
                    raw = {"기본": raw}
                self.data = raw
            except (OSError, ValueError):
                self.data = {}

    def get(self, printer, tcid):
        return dict(EMPTY, **self.data.get(printer, {}).get(tcid, {}))

    def set(self, printer, tcid, **kw):
        r = self.get(printer, tcid)
        r.update({k: v for k, v in kw.items() if v is not None})
        if "result" in kw and not kw.get("date"):
            r["date"] = datetime.date.today().isoformat()
        self.data.setdefault(printer, {})[tcid] = r
        self.save()

    def add_note(self, printer, tcid, line, **kw):
        r = self.get(printer, tcid)
        note = (r["note"] + "\n" if r["note"] else "") + line
        self.set(printer, tcid, note=note, **kw)

    def save(self):
        os.makedirs(os.path.dirname(self.path), exist_ok=True)
        tmp = self.path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.data, f, ensure_ascii=False, indent=1)
        os.replace(tmp, self.path)

    def printers(self):
        return list(self.data)

    def counts(self, printer):
        c = {"Pass": 0, "Fail": 0, "N/A": 0, "미실시": 0}
        for t in TC:
            res = self.get(printer, t["id"])["result"]
            c[res] = c.get(res, 0) + 1
        return c

    def export_csv(self, path=None):
        os.makedirs(RESULT_DIR, exist_ok=True)
        path = path or os.path.join(RESULT_DIR, f"결과_{datetime.datetime.now():%Y%m%d_%H%M%S}.csv")
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["프린터", "구분", "ID", "항목", "합격 기준", "결과", "시험자", "시험일", "FW 버전", "결함 ID", "비고"])
            for printer in self.printers():
                for t in TC:
                    r = self.get(printer, t["id"])
                    w.writerow([printer, t["cat"], t["id"], t["name"], t["criterion"], r["result"], r["tester"],
                                r["date"], r["fw"], r["defect"], r["note"]])
        return path

    def export_xlsx(self, plan_path, printer, out_path=None):
        """검증 계획표 엑셀의 '체크리스트' 시트에 해당 프린터 결과를 채워 새 파일로 저장."""
        from openpyxl import load_workbook
        wb = load_workbook(plan_path)
        if "체크리스트" not in wb.sheetnames:
            raise ValueError("'체크리스트' 시트가 없습니다. 검증 계획표 파일을 선택하세요.")
        ws = wb["체크리스트"]
        mine = self.data.get(printer, {})
        written = 0
        for row in range(1, ws.max_row + 1):
            tcid = ws.cell(row, COL_ID).value
            if tcid in mine:
                r = self.get(printer, tcid)
                ws.cell(row, COL_RESULT).value = r["result"]
                ws.cell(row, COL_TESTER).value = r["tester"]
                ws.cell(row, COL_DATE).value = r["date"]
                ws.cell(row, COL_FW).value = r["fw"]
                ws.cell(row, COL_DEFECT).value = r["defect"]
                ws.cell(row, COL_NOTE).value = r["note"]
                written += 1
        wb.calculation.fullCalcOnLoad = True
        if not out_path:
            root, ext = os.path.splitext(plan_path)
            out_path = f"{root}_결과_{_safe(printer)}_{datetime.datetime.now():%m%d_%H%M}{ext}"
        wb.save(out_path)
        return out_path, written
